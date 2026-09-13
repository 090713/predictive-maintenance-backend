"""
Report routes - async report generation and download.
"""
import asyncio
import csv
import io
import os
from datetime import datetime, timezone
from typing import Optional, List, Dict, Any
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, status, Query, BackgroundTasks
from fastapi.responses import StreamingResponse
from pymongo import ASCENDING, DESCENDING
from bson import ObjectId

from backend.config import settings
from backend.models.report import ReportCreate, ReportResponse, ReportListResponse, ReportListQuery, ReportStatus, ReportFormat
from backend.models.user import UserResponse, UserRole
from backend.services.auth import get_current_user_response, require_supervisor_or_admin
from backend.db.mongodb import get_collection


router = APIRouter(prefix="/api/v1/reports", tags=["Reports"])

# Reports storage directory
REPORTS_DIR = Path("reports")
REPORTS_DIR.mkdir(exist_ok=True)


async def generate_report_file(report_id: str):
    """Background task to generate report file."""
    reports_collection = get_collection("reports")
    assessments_collection = get_collection("assessments")

    # Get report
    doc = await reports_collection.find_one({"_id": ObjectId(report_id)})
    if not doc:
        return

    try:
        # Update status to generating
        await reports_collection.update_one(
            {"_id": ObjectId(report_id)},
            {"$set": {"status": "generating", "updated_at": datetime.now(timezone.utc)}}
        )

        # Build assessment query
        query = {}
        if doc["mission_ids"]:
            query["mission_id"] = {"$in": doc["mission_ids"]}
        if doc["machine_ids"]:
            query["machine_id"] = {"$in": doc["machine_ids"]}
        if doc["date_from"] or doc["date_to"]:
            date_query = {}
            if doc["date_from"]:
                date_query["$gte"] = doc["date_from"]
            if doc["date_to"]:
                date_query["$lte"] = doc["date_to"]
            query["ts"] = date_query

        # Get assessments
        cursor = get_collection("assessments").find(query).sort("ts", 1)
        assessments = []
        async for doc_assess in cursor:
            assessments.append(doc_assess)

        # Generate file
        timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        filename = f"report_{report_id}_{timestamp}.{doc['format']}"
        filepath = REPORTS_DIR / filename

        if doc["format"] == "csv":
            await write_csv_report(filepath, assessments, doc)
        else:
            await write_xlsx_report(filepath, assessments, doc)

        # Update report with file info
        file_size = filepath.stat().st_size
        await reports_collection.update_one(
            {"_id": ObjectId(report_id)},
            {"$set": {
                "status": "ready",
                "file_path": str(filepath),
                "file_size": file_size,
                "record_count": len(assessments),
                "completed_at": datetime.now(timezone.utc),
                "updated_at": datetime.now(timezone.utc),
            }}
        )
    except Exception as e:
        await reports_collection.update_one(
            {"_id": ObjectId(report_id)},
            {"$set": {
                "status": "failed",
                "error_message": str(e),
                "updated_at": datetime.now(timezone.utc),
            }}
        )


async def write_csv_report(filepath: Path, assessments: List[dict], report_doc: dict):
    """Write assessments to CSV file."""
    if not assessments:
        # Empty CSV with headers
        with open(filepath, "w", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(get_csv_headers(report_doc))
        return

    with open(filepath, "w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(get_csv_headers(report_doc))

        for assess in assessments:
            pred = assess.get("prediction", {})
            inputs = assess.get("inputs", {})
            row = [
                assess.get("machine_id", ""),
                assess.get("mission_id", ""),
                assess.get("user_id", ""),
                assess.get("ts", ""),
                pred.get("failure_probability", 0),
                pred.get("health_status", ""),
                pred.get("risk_level", ""),
                pred.get("predicted_class", ""),
                pred.get("failure_mode", ""),
                pred.get("failure_mode_name", ""),
                pred.get("anomaly_percentile", 0),
                pred.get("anomaly_score", 0),
                pred.get("source", ""),
                pred.get("model_version", ""),
                pred.get("prediction_id", ""),
                pred.get("latency_ms", 0),
            ]
            if report_doc.get("include_derived"):
                derived = pred.get("derived_params", {})
                row.extend([
                    derived.get("temp_difference", 0),
                    derived.get("mechanical_power", 0),
                    derived.get("overstrain", 0),
                ])
            if report_doc.get("include_features"):
                features = pred.get("contributing_features", [])
                for feat in features[:5]:  # Top 5 features
                    row.extend([feat.get("label", ""), feat.get("value", 0), feat.get("direction", "")])
            writer.writerow(row)


def get_csv_headers(report_doc: dict) -> List[str]:
    """Get CSV headers based on report options."""
    headers = [
        "machine_id", "mission_id", "user_id", "timestamp",
        "failure_probability", "health_status", "risk_level",
        "predicted_class", "failure_mode", "failure_mode_name",
        "anomaly_percentile", "anomaly_score", "source",
        "model_version", "prediction_id", "latency_ms"
    ]
    if report_doc.get("include_derived"):
        headers.extend(["temp_difference", "mechanical_power", "overstrain"])
    if report_doc.get("include_features"):
        for i in range(1, 6):
            headers.extend([f"feature_{i}_label", f"feature_{i}_value", f"feature_{i}_direction"])
    return headers


async def write_xlsx_report(filepath: Path, assessments: List[dict], report_doc: dict):
    """Write assessments to XLSX file."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Assessments"

    # Header row
    headers = get_csv_headers(report_doc)
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header.replace("_", " ").title())
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")

    for row_idx, assess in enumerate(assessments, 2):
        pred = assess.get("prediction", {})
        inputs = assess.get("inputs", {})
        row_data = [
            assess.get("machine_id", ""),
            assess.get("mission_id", ""),
            assess.get("user_id", ""),
            assess.get("ts", ""),
            pred.get("failure_probability", 0),
            pred.get("health_status", ""),
            pred.get("risk_level", ""),
            pred.get("predicted_class", ""),
            pred.get("failure_mode", ""),
            pred.get("failure_mode_name", ""),
            pred.get("anomaly_percentile", 0),
            pred.get("anomaly_score", 0),
            pred.get("source", ""),
            pred.get("model_version", ""),
            pred.get("prediction_id", ""),
            pred.get("latency_ms", 0),
        ]
        if report_doc.get("include_derived"):
            derived = pred.get("derived_params", {})
            row_data.extend([
                derived.get("temp_difference", 0),
                derived.get("mechanical_power", 0),
                derived.get("overstrain", 0),
            ])
        if report_doc.get("include_features"):
            features = pred.get("contributing_features", [])
            for feat in features[:5]:
                row_data.extend([feat.get("label", ""), feat.get("value", 0), feat.get("direction", "")])

        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    wb.save(filepath)


router = APIRouter(prefix="/api/v1/reports", tags=["Reports"])


@router.post("", response_model=ReportResponse, status_code=status.HTTP_201_CREATED)
async def create_report(
    report_data: ReportCreate,
    background_tasks: BackgroundTasks,
    current_user: UserResponse = Depends(require_supervisor_or_admin)
):
    """Create a new report generation job."""
    reports_collection = get_collection("reports")

    # Validate missions exist
    if report_data.mission_ids:
        missions_collection = get_collection("missions")
        count = await missions_collection.count_documents({"mission_id": {"$in": report_data.mission_ids}})
        if count != len(report_data.mission_ids):
            raise HTTPException(status_code=400, detail="One or more mission_ids do not exist")

    # Validate machines exist
    if report_data.machine_ids:
        machines_collection = get_collection("machines")
        count = await machines_collection.count_documents({"machine_id": {"$in": report_data.machine_ids}})
        if count != len(report_data.machine_ids):
            raise HTTPException(status_code=400, detail="One or more machine_ids do not exist")

    # Validate date range
    if report_data.date_from >= report_data.date_to:
        raise HTTPException(status_code=400, detail="date_from must be before date_to")

    report_doc = report_data.model_dump()
    report_doc["requested_by"] = current_user.id
    report_doc["status"] = "pending"
    report_doc["created_at"] = datetime.now(timezone.utc)
    report_doc["updated_at"] = datetime.now(timezone.utc)

    result = await get_collection("reports").insert_one(report_doc)
    report_doc["_id"] = str(result.inserted_id)

    # Schedule background generation
    background_tasks.add_task(generate_report_file, report_doc["_id"])

    return ReportResponse.from_db(report_doc)


@router.get("", response_model=ReportListResponse)
async def list_reports(
    query: ReportListQuery = Depends(),
    current_user: UserResponse = Depends(get_current_user_response)
):
    """List reports with filters and pagination."""
    reports_collection = get_collection("reports")

    # Build query based on role
    query = {}
    if current_user.role == "worker":
        query["requested_by"] = current_user.id

    if query.status:
        query["status"] = query.status.value
    if query.requested_by and current_user.role in ("admin", "supervisor"):
        query["requested_by"] = query.requested_by
    if query.date_from or query.date_to:
        date_query = {}
        if query.date_from:
            date_query["$gte"] = query.date_from
        if query.date_to:
            date_query["$lte"] = query.date_to
        query["created_at"] = date_query

    total = await reports_collection.count_documents(query)
    total_pages = (total + query.page_size - 1) // query.page_size

    cursor = reports_collection.find(query).sort("created_at", -1).skip((query.page - 1) * query.page_size).limit(query.page_size)

    items = []
    async for doc in cursor:
        doc["_id"] = str(doc["_id"])
        items.append(ReportResponse.from_db(doc))

    return ReportListResponse(
        items=items,
        total=total,
        page=query.page,
        page_size=query.page_size,
        total_pages=total_pages,
    )


@router.get("/{report_id}", response_model=ReportResponse)
async def get_report(
    report_id: str,
    current_user: UserResponse = Depends(get_current_user_response)
):
    """Get a single report by ID."""
    doc = await get_collection("reports").find_one({"_id": ObjectId(report_id)})
    if not doc:
        raise HTTPException(status_code=404, detail="Report not found")

    if current_user.role == "worker" and doc["requested_by"] != current_user.id:
        raise HTTPException(status_code=403, detail="Access denied")

    doc["_id"] = str(doc["_id"])
    return ReportResponse.from_db(doc)


@router.get("/{report_id}/download")
async def download_report(
    report_id: str,
    current_user: UserResponse = Depends(get_current_user_response)
):
    """Download a generated report file."""
    doc = await get_collection("reports").find_one({"_id": ObjectId(report_id)})
    if not doc:
        raise HTTPException(status_code=404, detail="Report not found")

    if current_user.role == "worker" and doc["requested_by"] != current_user.id:
        raise HTTPException(status_code=403, detail="Access denied")

    if doc["status"] != "ready" or not doc.get("file_path"):
        raise HTTPException(status_code=400, detail="Report not ready for download")

    filepath = Path(doc["file_path"])
    if not filepath.exists():
        raise HTTPException(status_code=404, detail="Report file not found")

    media_type = "text/csv" if doc["format"] == "csv" else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    filename = filepath.name

    return StreamingResponse(
        open(filepath, "rb"),
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


@router.delete("/{report_id}")
async def delete_report(
    report_id: str,
    current_user: UserResponse = Depends(require_supervisor_or_admin)
):
    """Delete a report and its file."""
    reports_collection = get_collection("reports")

    doc = await reports_collection.find_one({"_id": ObjectId(report_id)})
    if not doc:
        raise HTTPException(status_code=404, detail="Report not found")

    # Delete file if exists
    if doc.get("file_path"):
        filepath = Path(doc["file_path"])
        if filepath.exists():
            filepath.unlink()

    await reports_collection.delete_one({"_id": ObjectId(report_id)})
    return {"message": "Report deleted successfully"}